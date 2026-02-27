import json
from io import BytesIO
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods, require_POST
from django.utils import timezone
from .models import IdeaGenerationRequest, BlogDraft
from .services import (
    generate_article_ideas, generate_blog_draft, generate_blog_outline,
    generate_blog_from_outline, generate_trending_ideas, generate_blog_image,
    download_and_save_image,
    MODEL_CLAUDE, MODEL_OPENAI, MODEL_GEMINI, MODEL_AUTO,
)
from articles.models import Article, Category


# ---------------------------------------------------------------------------
# Topic choices — grouped by category for the idea generator
# ---------------------------------------------------------------------------

TOPIC_CHOICES = [
    ('people',      'People',               'bi-person-fill'),
    ('events',      'Events',               'bi-calendar-event-fill'),
    ('places',      'Places',               'bi-geo-alt-fill'),
    ('culture',     'Culture',              'bi-palette-fill'),
    ('technology',  'Technology',           'bi-cpu-fill'),
    ('lifestyle',   'Lifestyle',            'bi-heart-fill'),
    ('business',    'Business',             'bi-briefcase-fill'),
    ('health',      'Health & Wellness',    'bi-heart-pulse-fill'),
    ('travel',      'Travel',               'bi-airplane-fill'),
    ('food',        'Food & Cuisine',       'bi-egg-fried'),
    ('arts',        'Arts & Entertainment', 'bi-music-note-beamed'),
    ('sports',      'Sports',               'bi-trophy-fill'),
    ('environment', 'Environment',          'bi-tree-fill'),
    ('science',     'Science',              'bi-flask-fill'),
]

TOPIC_CATEGORIES = [
    {
        'name': 'People & Culture',
        'icon': 'bi-people-fill',
        'topics': [
            ('people',    'People',    'bi-person-fill'),
            ('culture',   'Culture',   'bi-palette-fill'),
            ('lifestyle', 'Lifestyle', 'bi-heart-fill'),
            ('arts',      'Arts & Entertainment', 'bi-music-note-beamed'),
        ],
    },
    {
        'name': 'World & Travel',
        'icon': 'bi-globe',
        'topics': [
            ('events',      'Events',      'bi-calendar-event-fill'),
            ('places',      'Places',      'bi-geo-alt-fill'),
            ('travel',      'Travel',      'bi-airplane-fill'),
            ('environment', 'Environment', 'bi-tree-fill'),
        ],
    },
    {
        'name': 'Health & Lifestyle',
        'icon': 'bi-heart-pulse-fill',
        'topics': [
            ('health',  'Health & Wellness', 'bi-heart-pulse-fill'),
            ('food',    'Food & Cuisine',    'bi-egg-fried'),
            ('sports',  'Sports',            'bi-trophy-fill'),
        ],
    },
    {
        'name': 'Business & Tech',
        'icon': 'bi-briefcase-fill',
        'topics': [
            ('business',    'Business',    'bi-briefcase-fill'),
            ('technology',  'Technology',  'bi-cpu-fill'),
            ('science',     'Science',     'bi-flask-fill'),
        ],
    },
]

VALID_TOPICS = {t[0] for t in TOPIC_CHOICES}
VALID_MODELS = {MODEL_CLAUDE, MODEL_OPENAI, MODEL_GEMINI, MODEL_AUTO}


# ---------------------------------------------------------------------------
# Idea Generator
# ---------------------------------------------------------------------------

@login_required
def idea_generator_view(request):
    if request.user.is_staff:
        recent_requests = IdeaGenerationRequest.objects.select_related('user').order_by('-created_at')[:10]
    else:
        recent_requests = IdeaGenerationRequest.objects.filter(user=request.user)[:5]

    preselected_topics = []
    topics_param = request.GET.get('topics', '')
    if topics_param:
        preselected_topics = [t.strip() for t in topics_param.split(',') if t.strip() in VALID_TOPICS]

    return render(request, 'ai_tools/idea_generator.html', {
        'topic_categories': TOPIC_CATEGORIES,
        'topic_choices': TOPIC_CHOICES,
        'recent_requests': recent_requests,
        'preselected_topics': json.dumps(preselected_topics),
        'is_admin': request.user.is_staff,
    })


@login_required
@require_POST
def generate_ideas_ajax(request):
    try:
        data = json.loads(request.body)
        topics = data.get('topics', [])
        additional_context = data.get('context', '').strip()
        model = data.get('model', MODEL_CLAUDE)

        if not topics:
            return JsonResponse({'error': 'Please select at least one topic.'}, status=400)
        if len(topics) > 8:
            return JsonResponse({'error': 'Please select at most 8 topics.'}, status=400)
        for topic in topics:
            if topic not in VALID_TOPICS:
                return JsonResponse({'error': f'Invalid topic: {topic}'}, status=400)
        if additional_context and len(additional_context) > 500:
            return JsonResponse({'error': 'Additional context is too long (max 500 chars).'}, status=400)
        if model not in VALID_MODELS:
            model = MODEL_CLAUDE

        ideas = generate_article_ideas(topics, additional_context, model=model)

        IdeaGenerationRequest.objects.create(
            user=request.user,
            topics=topics,
            additional_context=additional_context,
            generated_ideas=ideas,
            ai_model_used=model,
        )

        return JsonResponse({'ideas': ideas, 'model_used': model})

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=503)
    except Exception:
        return JsonResponse({'error': 'Failed to generate ideas. Please try again.'}, status=500)


# ---------------------------------------------------------------------------
# Blog Generator — Main Page
# ---------------------------------------------------------------------------

@login_required
def blog_generator_view(request):
    if request.user.is_staff:
        drafts = BlogDraft.objects.select_related('user').order_by('-created_at')[:10]
    else:
        drafts = BlogDraft.objects.filter(user=request.user).order_by('-created_at')[:10]

    return render(request, 'ai_tools/blog_generator.html', {
        'drafts': drafts,
        'is_admin': request.user.is_staff,
    })


# ---------------------------------------------------------------------------
# Blog Generator — Step 1: Generate Outline
# ---------------------------------------------------------------------------

@login_required
@require_POST
def generate_outline_ajax(request):
    """Step 1: Generate a SEO-friendly blog outline from a topic."""
    try:
        data = json.loads(request.body)
        topic = data.get('topic', '').strip()
        model = data.get('model', MODEL_CLAUDE)

        if not topic:
            return JsonResponse({'error': 'Please enter a topic or idea.'}, status=400)
        if len(topic) < 5:
            return JsonResponse({'error': 'Topic too short — please be more specific.'}, status=400)
        if len(topic) > 500:
            return JsonResponse({'error': 'Topic too long (max 500 characters).'}, status=400)
        if model not in VALID_MODELS:
            model = MODEL_CLAUDE

        outline = generate_blog_outline(topic, model=model)
        return JsonResponse({'outline': outline, 'model_used': model})

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=503)
    except Exception as e:
        return JsonResponse({'error': f'Failed to generate outline: {str(e)}'}, status=500)


# ---------------------------------------------------------------------------
# Blog Generator — Step 2: Generate Full Blog from Outline
# ---------------------------------------------------------------------------

@login_required
@require_POST
def generate_draft_ajax(request):
    """Step 2: Generate full blog from outline, or directly from idea (fallback)."""
    try:
        data = json.loads(request.body)
        idea = data.get('idea', '').strip()
        model = data.get('model', MODEL_CLAUDE)
        outline = data.get('outline', {})

        if not idea:
            return JsonResponse({'error': 'Please provide an idea or topic.'}, status=400)
        if len(idea) < 5:
            return JsonResponse({'error': 'Idea too short — please provide more detail.'}, status=400)
        if len(idea) > 1000:
            return JsonResponse({'error': 'Idea too long (max 1000 characters).'}, status=400)
        if model not in VALID_MODELS:
            model = MODEL_CLAUDE

        if outline and outline.get('sections'):
            result = generate_blog_from_outline(idea, outline, model=model)
        else:
            result = generate_blog_draft(idea, model=model)

        seo_keywords = outline.get('focus_keywords', []) if outline else []
        summary_bullets = result.get('summary_bullets', [])

        draft = BlogDraft.objects.create(
            user=request.user,
            idea=idea,
            title=result.get('title', ''),
            generated_content=result.get('content', ''),
            edited_content=result.get('content', ''),
            status='draft',
            ai_model_used=model,
            outline=outline or {},
            summary=summary_bullets,
            seo_keywords=seo_keywords,
        )

        return JsonResponse({
            'draft_id': draft.id,
            'title': draft.title,
            'excerpt': result.get('excerpt', ''),
            'content': draft.generated_content,
            'summary_bullets': summary_bullets,
            'seo_keywords': seo_keywords,
            'model_used': model,
        })

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=503)
    except Exception as e:
        return JsonResponse({'error': f'Failed to generate blog: {str(e)}'}, status=500)


# ---------------------------------------------------------------------------
# Blog Generator — Step 3 (optional): Generate Image
# ---------------------------------------------------------------------------

@login_required
@require_POST
def generate_image_ajax(request, draft_id):
    """Generate a featured image for a blog draft using DALL-E 3."""
    if request.user.is_staff:
        draft = get_object_or_404(BlogDraft, id=draft_id)
    else:
        draft = get_object_or_404(BlogDraft, id=draft_id, user=request.user)

    try:
        excerpt = draft.edited_content[:300] if draft.edited_content else ''
        image_url = generate_blog_image(draft.title or draft.idea, excerpt)

        # Download and save image to draft
        image_bytes = download_and_save_image(image_url, 'generated.jpg')
        from django.core.files.base import ContentFile
        draft.featured_image.save(
            f'ai_generated_{draft.id}.jpg',
            ContentFile(image_bytes),
            save=True,
        )

        return JsonResponse({'image_url': draft.featured_image.url})

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=503)
    except Exception as e:
        return JsonResponse({'error': f'Image generation failed: {str(e)}'}, status=500)


# ---------------------------------------------------------------------------
# Edit Draft
# ---------------------------------------------------------------------------

@login_required
def edit_draft_view(request, draft_id):
    if request.user.is_staff:
        draft = get_object_or_404(BlogDraft, id=draft_id)
    else:
        draft = get_object_or_404(BlogDraft, id=draft_id, user=request.user)

    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        content = request.POST.get('content', '').strip()
        action = request.POST.get('action', 'save')

        if not title:
            messages.error(request, 'Title is required.')
            return redirect('edit_draft', draft_id=draft.id)
        if not content:
            messages.error(request, 'Content is required.')
            return redirect('edit_draft', draft_id=draft.id)

        draft.title = title
        draft.edited_content = content

        if 'featured_image' in request.FILES:
            draft.featured_image = request.FILES['featured_image']

        if action == 'publish':
            category_id = request.POST.get('category')
            category = None
            if category_id:
                try:
                    category = Category.objects.get(id=category_id)
                except Category.DoesNotExist:
                    pass

            article, created = (
                (draft.linked_article, False)
                if draft.linked_article
                else (None, False)
            )

            if article:
                # Update existing article and publish it
                article.title = title
                article.content = content
                article.excerpt = request.POST.get('excerpt', article.excerpt)[:500]
                article.category = category or article.category
                article.status = 'published'
                article.published_at = timezone.now()
                article.ai_generated = True
                if draft.featured_image:
                    article.cover_image = draft.featured_image
                article.save()
            else:
                article = Article.objects.create(
                    title=title,
                    content=content,
                    excerpt=request.POST.get('excerpt', '')[:500],
                    author=request.user,
                    category=category,
                    status='published',
                    published_at=timezone.now(),
                    ai_generated=True,
                )
                if draft.featured_image:
                    article.cover_image = draft.featured_image
                    article.save()

            draft.status = 'published'
            draft.linked_article = article
            draft.save()
            messages.success(request, f'Blog published successfully! "{title}"')
            return redirect('article_detail', slug=article.slug)

        else:  # save
            draft.status = 'saved'
            draft.save()
            messages.success(request, 'Draft saved successfully.')
            return redirect('edit_draft', draft_id=draft.id)

    return render(request, 'ai_tools/edit_draft.html', {
        'draft': draft,
        'categories': Category.objects.all(),
    })


# ---------------------------------------------------------------------------
# Delete Draft
# ---------------------------------------------------------------------------

@login_required
@require_POST
def delete_draft_view(request, draft_id):
    if request.user.is_staff:
        draft = get_object_or_404(BlogDraft, id=draft_id)
    else:
        draft = get_object_or_404(BlogDraft, id=draft_id, user=request.user)
    draft.delete()
    messages.success(request, 'Draft deleted.')
    return redirect('manage_blogs')


# ---------------------------------------------------------------------------
# Manage Blogs & Ideas
# ---------------------------------------------------------------------------

@login_required
def manage_blogs_view(request):
    if request.user.is_staff:
        drafts = BlogDraft.objects.select_related('user', 'linked_article').order_by('-created_at')
    else:
        drafts = BlogDraft.objects.filter(user=request.user).select_related('linked_article').order_by('-created_at')

    return render(request, 'ai_tools/manage_blogs.html', {
        'drafts': drafts,
        'is_admin': request.user.is_staff,
    })


@login_required
def manage_ideas_view(request):
    if request.user.is_staff:
        idea_requests = IdeaGenerationRequest.objects.select_related('user').order_by('-created_at')
    else:
        idea_requests = IdeaGenerationRequest.objects.filter(user=request.user).order_by('-created_at')

    return render(request, 'ai_tools/manage_ideas.html', {
        'idea_requests': idea_requests,
        'is_admin': request.user.is_staff,
    })


# ---------------------------------------------------------------------------
# Trending Ideas (Admin Only)
# ---------------------------------------------------------------------------

@login_required
def trending_ideas_view(request):
    if not request.user.is_staff:
        messages.error(request, 'Admin access required.')
        return redirect('dashboard')

    return render(request, 'ai_tools/trending_ideas.html', {
        'topic_categories': TOPIC_CATEGORIES,
    })


@login_required
@require_POST
def generate_trending_ajax(request):
    if not request.user.is_staff:
        return JsonResponse({'error': 'Admin access required.'}, status=403)
    try:
        data = json.loads(request.body)
        category = data.get('category', '').strip()
        model = data.get('model', MODEL_CLAUDE)
        if model not in VALID_MODELS:
            model = MODEL_CLAUDE

        ideas = generate_trending_ideas(category, model=model)
        return JsonResponse({'ideas': ideas, 'model_used': model})

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=503)
    except Exception as e:
        return JsonResponse({'error': f'Failed to generate trending topics: {str(e)}'}, status=500)


# ---------------------------------------------------------------------------
# Download Draft as DOCX
# ---------------------------------------------------------------------------

@login_required
def download_draft_docx(request, draft_id):
    if request.user.is_staff:
        draft = get_object_or_404(BlogDraft, id=draft_id)
    else:
        draft = get_object_or_404(BlogDraft, id=draft_id, user=request.user)

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        import re as _re

        doc = Document()
        # Title
        title_para = doc.add_heading(draft.title or 'Blog Post', 0)

        # Meta info
        doc.add_paragraph(f'Author: {draft.user.get_full_name() or draft.user.username}')
        doc.add_paragraph(f'Generated: {draft.created_at.strftime("%B %d, %Y")}')
        doc.add_paragraph(f'Model: {draft.ai_model_used.upper()}')

        if draft.seo_keywords:
            doc.add_paragraph(f'SEO Keywords: {", ".join(draft.seo_keywords)}')

        doc.add_paragraph('')  # spacer

        # Content
        content = draft.edited_content or draft.generated_content
        for line in content.split('\n'):
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith('### '):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith('## '):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith('# '):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith(('- ', '* ', '• ')):
                clean = _re.sub(r'\*\*(.*?)\*\*', r'\1', stripped[2:])
                clean = _re.sub(r'\*(.*?)\*', r'\1', clean)
                doc.add_paragraph(clean, style='List Bullet')
            else:
                clean = _re.sub(r'\*\*(.*?)\*\*', r'\1', stripped)
                clean = _re.sub(r'\*(.*?)\*', r'\1', clean)
                doc.add_paragraph(clean)

        # Summary bullets
        if draft.summary:
            doc.add_heading('Key Takeaways', level=2)
            bullets = draft.summary if isinstance(draft.summary, list) else [draft.summary]
            for b in bullets:
                if b.strip():
                    doc.add_paragraph(b.strip('•- '), style='List Bullet')

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = (draft.title or 'blog').replace(' ', '_')[:60] + '.docx'
        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except ImportError:
        messages.error(request, 'python-docx is not installed. Run: pip install python-docx')
        return redirect('edit_draft', draft_id=draft.id)
    except Exception as e:
        messages.error(request, f'DOCX export failed: {str(e)}')
        return redirect('edit_draft', draft_id=draft.id)


# ---------------------------------------------------------------------------
# Download Draft as PDF (print-friendly page, user saves via browser)
# ---------------------------------------------------------------------------

@login_required
def download_draft_pdf(request, draft_id):
    """Render a print-optimised HTML page — user presses Ctrl+P to save as PDF."""
    if request.user.is_staff:
        draft = get_object_or_404(BlogDraft, id=draft_id)
    else:
        draft = get_object_or_404(BlogDraft, id=draft_id, user=request.user)

    return render(request, 'ai_tools/blog_print.html', {'draft': draft})
